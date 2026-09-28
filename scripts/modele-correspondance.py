#!/usr/bin/env python3
"""Modèle vierge de la table de correspondance identifiant ↔ agent.

Décision du 28/09/2026 (question 83, choix a) : la correspondance entre
l'identifiant généré par le site (AG-001…) et la personne est un classeur
tenu hors du site, sur le réseau de l'établissement, dans un dossier à droits
restreints (question 28, choix a). Ce script n'écrit que des en-têtes, des
contrôles et un mode d'emploi : aucun nom, le dépôt étant public.

    python3 scripts/modele-correspondance.py        # openpyxl requis

Sortie : docs/modeles/table-correspondance-agents.xlsx. Pour changer le
modèle, modifier ce script et le relancer ; test/modele-correspondance.test.ts
vérifie que le classeur livré reste vierge.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SORTIE = Path(__file__).resolve().parent.parent / "docs" / "modeles" / "table-correspondance-agents.xlsx"

# Garde-fous actifs de la ligne 3 à la ligne 2000 : des décennies d'arrivées
# pour une unité de quelques dizaines d'agents.
PREMIERE, DERNIERE = 3, 2000

# (en-tête, largeur, colonne de dates) — question 83, choix a.
COLONNES = [
    ("Identifiant", 13, False),
    ("Nom", 24, False),
    ("Prénom", 18, False),
    ("Fonction", 22, False),
    ("Identifiant créé le", 19, True),
    ("Créé par", 26, False),
    ("Remis à l'agent le", 19, True),
    ("Clos le", 13, True),
]

POLICE = "Arial"
LETTRES = "ABCDEFGH"


def identifiant_valide(cellule: str) -> str:
    """Vrai si la cellule porte un identifiant au format du site, sans doublon.

    Format de lib/identifiant.ts : « AG- » puis trois chiffres au moins, sans
    zéro de tête au-delà de trois (AG-001, AG-017, AG-1000) ; AG-000 n'est
    jamais généré. La comparaison au nombre reformaté écarte tout ce qui n'est
    pas une suite de chiffres (espace, signe, virgule, exposant).
    """
    chiffres = f"MID({cellule},4,20)"
    return (
        "IFERROR(AND("
        f'EXACT(LEFT({cellule},3),"AG-"),'
        f"LEN({cellule})>=6,"
        f'{chiffres}=TEXT(VALUE({chiffres}),REPT("0",LEN({cellule})-3)),'
        f"VALUE({chiffres})>=1,"
        f'OR(LEN({cellule})=6,LEFT({chiffres},1)<>"0"),'
        f"COUNTIF($A${PREMIERE}:$A${DERNIERE},{cellule})=1"
        "),FALSE)"
    )


MODE_EMPLOI = [
    ("À quoi elle sert",
     "Le site ne connaît les agents que par leur identifiant (AG-001, AG-002…) : les rapports, "
     "le registre et le répertoire qu'il exporte n'en portent pas d'autre. Cette table est le seul "
     "endroit où l'identifiant rejoint la personne : elle sert à retrouver qui est derrière un "
     "identifiant, et à porter le nom sur l'édition d'un rapport. Au sens du RGPD, c'est "
     "l'« information supplémentaire » de la pseudonymisation, à conserver séparément et à "
     "protéger (art. 4 § 5)."),
    ("Où elle est rangée",
     "Sur le réseau de l'établissement, dans un dossier à droits restreints, lisible par les "
     "tuteurs qui visent les rapports : [à compléter : chemin du dossier ; qui attribue les "
     "droits]. Une seule table : pas de copie de travail ailleurs."),
    ("Qui y écrit",
     "Le pharmacien responsable, qui la tient. Les tuteurs créent aussi des identifiants : "
     "[à préciser] s'ils écrivent eux-mêmes la ligne ou la font écrire par le pharmacien."),
    ("Où saisir",
     "Feuille « Correspondance » : une ligne par identifiant, à partir de la ligne 3. Les lignes 1 "
     "et 2 (titre et en-têtes) ne se modifient pas."),
    ("À la création d'un identifiant",
     "Sur le site, écran Personnel, « Créer un identifiant ». Ajouter la ligne aussitôt : "
     "l'identifiant tel que le site l'affiche, le nom, le prénom, la fonction, la date de création "
     "et qui l'a créé. Remettre ensuite l'identifiant à l'agent, avec l'information prévue "
     "[à compléter : note de service], et noter la date dans « Remis à l'agent le »."),
    ("À l'édition d'un rapport",
     "Chercher l'identifiant (Ctrl+F, ou le filtre de la colonne Identifiant), puis reporter le "
     "nom, le prénom et la fonction dans « Éditer avec le nom de l'agent », sur la fiche du "
     "rapport."),
    ("Au départ de l'agent",
     "Clore l'identifiant sur le site (écran Personnel, « Clore ») et noter la date dans « Clos "
     "le ». La ligne reste."),
    ("Conservation",
     "Une ligne se conserve aussi longtemps que le dossier d'habilitation de l'agent et s'efface "
     "quand ce dossier est détruit (durée [à vérifier])."),
    ("Contrôle",
     "Comparer la liste des identifiants de l'écran Personnel à cette table : chaque identifiant "
     "du site doit y avoir sa ligne. Périodicité [à préciser]."),
    ("Jamais",
     "L'envoyer par courriel ou messagerie. En faire une copie hors du dossier : sur papier, sur "
     "une clé USB, sur un poste personnel, dans le dossier « Téléchargements ». La déposer sur le "
     "site, ou écrire un nom sur le site ailleurs que dans le formulaire d'édition du rapport "
     "(motif d'arbitrage, signalement…). La coller dans un outil en ligne, intelligence "
     "artificielle comprise."),
    ("Ce qu'elle ne contient pas",
     "Ni matricule, ni date de naissance, ni adresse, ni commentaire libre : le nom et le prénom "
     "suffisent à retrouver la personne (minimisation des données, RGPD art. 5 § 1 c)."),
    ("Garde-fous",
     "La colonne Identifiant refuse une saisie hors du format du site (AG- suivi d'au moins trois "
     "chiffres, AG-000 exclu) et un identifiant déjà présent. Une valeur collée échappe à ce "
     f"contrôle : la case vire alors au rouge. Contrôles actifs jusqu'à la ligne {DERNIERE}."),
    ("Droits des agents",
     "Un agent peut demander à voir sa ligne ou à la faire rectifier, auprès du pharmacien "
     "responsable ou du DPO."),
    ("Sauvegarde",
     "Celle du dossier réseau [à vérifier auprès du service informatique]. Perdue, la table ne se "
     "reconstitue qu'en rouvrant les dossiers d'habilitation un à un."),
]

EXEMPLE = [
    ("Identifiant", "AG-000"),
    ("Nom", "EXEMPLE"),
    ("Prénom", "Camille"),
    ("Fonction", "Préparateur"),
    ("Identifiant créé le", "28/09/2026"),
    ("Créé par", "EXEMPLE Dominique, tuteur"),
    ("Remis à l'agent le", "29/09/2026"),
    ("Clos le", "vide tant que l'agent est en poste"),
]

ORIGINE = (
    "Modèle vierge du 28/09/2026. Décisions du 18/09/2026 (questions 27 et 28) et du 28/09/2026 "
    "(question 83) du projet PharmaTechX, consignées dans docs/DECISIONS.md ; classeur fabriqué "
    "par scripts/modele-correspondance.py."
)


def police(**options) -> Font:
    return Font(name=POLICE, size=options.pop("size", 10), **options)


def feuille_correspondance(ws) -> None:
    ws.title = "Correspondance"
    for lettre, (_, largeur, date) in zip(LETTRES, COLONNES):
        colonne = ws.column_dimensions[lettre]
        colonne.width = largeur
        # Style de colonne : une cellule saisie plus tard le reprend.
        colonne.font = police()
        if date:
            colonne.number_format = "dd/mm/yyyy"

    ws.merge_cells("A1:H1")
    titre = ws["A1"]
    titre.value = (
        "Table de correspondance identifiant ↔ agent — confidentielle, ne sort jamais du dossier "
        "réseau protégé. Une ligne par identifiant, à partir de la ligne 3 ; mode d'emploi dans la "
        "feuille suivante."
    )
    titre.font = police(bold=True, color="9C0006")
    titre.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30

    trait = Side(style="thin", color="7F7F7F")
    for lettre, (entete, _, _) in zip(LETTRES, COLONNES):
        cellule = ws[f"{lettre}2"]
        cellule.value = entete
        cellule.font = police(bold=True)
        cellule.fill = PatternFill(fill_type="solid", start_color="D9E1F2", end_color="D9E1F2")
        cellule.border = Border(bottom=trait)
        cellule.alignment = Alignment(vertical="center")

    ws.freeze_panes = f"A{PREMIERE}"
    ws.sheet_view.selection[0].activeCell = ws.sheet_view.selection[0].sqref = f"A{PREMIERE}"
    ws.auto_filter.ref = f"A2:H{DERNIERE}"

    plage = f"A{PREMIERE}:A{DERNIERE}"
    controle = DataValidation(
        type="custom",
        formula1=identifiant_valide(f"A{PREMIERE}"),
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Identifiant refusé",
        error=(
            "Format du site attendu : AG- suivi d'au moins trois chiffres (AG-001, AG-017, AG-1000), "
            "AG-000 exclu. Un identifiant ne figure qu'une fois dans la table."
        ),
        showInputMessage=True,
        promptTitle="Identifiant",
        prompt="Tel que le site l'affiche sur l'écran Personnel, par exemple AG-017.",
    )
    controle.add(plage)
    ws.add_data_validation(controle)

    # Une valeur collée échappe au contrôle de saisie : elle vire au rouge.
    ws.conditional_formatting.add(
        plage,
        FormulaRule(
            formula=[f'AND(A{PREMIERE}<>"",NOT({identifiant_valide(f"A{PREMIERE}")}))'],
            fill=PatternFill(fill_type="solid", start_color="FFC7CE", end_color="FFC7CE"),
            font=Font(color="9C0006"),
        ),
    )


def feuille_mode_emploi(ws) -> None:
    ws.title = "Mode d'emploi"
    # Lignes d'environ 75 signes ; imprimée, la feuille tient en largeur sur un A4.
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 64
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    haut = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = "Table de correspondance identifiant ↔ agent — mode d'emploi"
    ws["A1"].font = police(bold=True, size=13)
    ws["A2"] = "Confidentiel : ce classeur ne sort jamais du dossier réseau protégé."
    ws["A2"].font = police(bold=True, color="9C0006")

    ligne = 4
    for rubrique, texte in MODE_EMPLOI:
        ws.cell(ligne, 1, rubrique).font = police(bold=True)
        ws.cell(ligne, 2, texte).font = police()
        ws.cell(ligne, 1).alignment = haut
        ws.cell(ligne, 2).alignment = haut
        ligne += 1

    ligne += 1
    ws.cell(ligne, 1, "Exemple fictif d'une ligne").font = police(bold=True)
    ws.cell(ligne, 2, "À ne pas recopier : AG-000 n'existe pas sur le site.").font = police(italic=True)
    ligne += 1
    for colonne, valeur in EXEMPLE:
        ws.cell(ligne, 1, colonne).font = police()
        ws.cell(ligne, 2, valeur).font = police(italic=True, color="595959")
        ligne += 1

    ligne += 1
    ws.cell(ligne, 1, "Origine").font = police(bold=True)
    ws.cell(ligne, 2, ORIGINE).font = police()
    ws.cell(ligne, 1).alignment = haut
    ws.cell(ligne, 2).alignment = haut


def main() -> None:
    classeur = Workbook()
    feuille_correspondance(classeur.active)
    feuille_mode_emploi(classeur.create_sheet())
    classeur.active = 0

    proprietes = classeur.properties
    proprietes.title = "Table de correspondance identifiant ↔ agent"
    proprietes.subject = "Modèle vierge — question 83, choix a"
    proprietes.creator = "PharmaTechX"
    proprietes.lastModifiedBy = "PharmaTechX"
    proprietes.created = proprietes.modified = datetime(2026, 9, 28)

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(SORTIE)
    print(SORTIE)


if __name__ == "__main__":
    main()
